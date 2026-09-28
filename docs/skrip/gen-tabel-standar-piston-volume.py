#!/usr/bin/env python3
"""
Reimplementasi rantai hitung DUA master piston volume (Fixed & Graduated),
diadu SEL DEMI SEL ke cache Excel. Dua mode:

  master : tiru persis (termasuk cacat G-2, G-8, G-9 dan sel suhu dari tautan luar)
  benar  : cacat salin-tempel yang MENGECILKAN U / merusak data dihitung benar
           - G-2  ambang tara-ulang m7 titik MIN 20 g   -> 200 g
           - G-8  m7 titik MID/MAX mengambil M4 (J32/L32) -> M7
           - G-9  Graduated P87 = SQRT(P86^2)+(H87^2)  -> SQRT(P86^2 + H87^2)
           - G-7  pita CMC bilangan bulat (1 < n < 2 "cek range") -> pita kontinu
  (Cacat-cacat G-2/G-8 cuma terpicu kalau massa > 20/200 g; di data contoh tidak.)

Keluaran (--tulis <repo>):
  database/data/tabel-standar-piston-volume.json
  database/data/sesi-master-piston-volume.json
"""
import datetime
import hashlib
import io
import json
import math
import os
import statistics as st
import sys

import msoffcrypto
import openpyxl
import mpmath as mp

PASSWORD = "spirit285"
TOL = 1e-12
mp.mp.dps = 40

FIXED = "Master_Olah_Data_Fixed_Piston_Volume.xlsm"
GRAD = "Master_Olah_Data_Graduated_Piston_Volume.xlsm"
NAMA_ASLI = {
    "fixed": "Master Olah Data_Fixed Piston Volume.xlsm",
    "graduated": "Master Olah Data_Graduated Piston Volume.xlsm",
}

A1, A2, A3, A4, A5 = -3.983035, 301.797, 522528.9, 69.34881, 999.97495
RHO_B = 8.0
ALFA = 0.0003
U_OPERATOR = 0.00015


def tinv(df):
    d = mp.mpf(int(df))
    f = lambda t: mp.betainc(d / 2, mp.mpf(1) / 2, 0, d / (d + t * t), regularized=True) / 2 - mp.mpf("0.025")
    return float(mp.findroot(f, 2.0))


def buka(path):
    buf = io.BytesIO()
    with open(path, "rb") as fh:
        off = msoffcrypto.OfficeFile(fh)
        off.load_key(password=PASSWORD)
        off.decrypt(buf)
    buf.seek(0)
    wf = openpyxl.load_workbook(buf)
    buf.seek(0)
    wv = openpyxl.load_workbook(buf, data_only=True)
    return wf, wv


def tgl(v):
    return v.strftime("%Y-%m-%d") if isinstance(v, (datetime.datetime, datetime.date)) else v


def tanaka(t):
    # Urutan operasi DISALIN dari sel: ($Q$57*(1-(((H52+$Q$53))^2*(H52+$Q$54))/($Q$55*(H52+$Q$56)))/1000)
    return A5 * (1 - (((t + A1)) ** 2 * (t + A2)) / (A3 * (t + A4))) / 1000


def rho_udara(P, RH, T):
    return ((0.34848 * P) - (0.009 * RH) * math.exp(0.061 * T)) / (T + 273.15) / 1000


def terdekat(daftar, x):
    best, jarak = None, None
    for v in daftar:
        d = abs(v - x)
        if jarak is None or d < jarak:
            best, jarak = v, d
    return best


# ---------------------------------------------------------------------------
# Tabel referensi (sama di kedua master — diadu)
# ---------------------------------------------------------------------------
def baca_tabel(path):
    _, wv = buka(path)
    db = wv["DATABASE"]
    timbangan = []
    for r in range(13, 16):
        timbangan.append({
            "nomor": db.cell(r, 21).value, "nama": db.cell(r, 22).value, "merk_tipe": db.cell(r, 23).value,
            "nomor_seri": str(db.cell(r, 24).value), "tertelusur": db.cell(r, 25).value,
            "jatuh_tempo": tgl(db.cell(r, 26).value), "u95": float(db.cell(r, 28).value),
            "resolusi": float(db.cell(r, 30).value), "stdev": float(db.cell(r, 31).value or 0),
        })
    termometer = {"nama": db["V16"].value, "merk_tipe": db["W16"].value, "nomor_seri": str(db["X16"].value),
                  "tertelusur": db["Y16"].value, "jatuh_tempo": tgl(db["Z16"].value), "u95": float(db["AB16"].value)}
    sensor = {"nama": "Sensor Suhu PRT PT100", "jatuh_tempo": tgl(db["Z17"].value), "u95": float(db["AB17"].value)}
    sk = wv["STANDARD KALIBRATOR"]
    meter = [{"indeks": float(sk.cell(r, 2).value), "koreksi": float(sk.cell(r, 3).value)}
             for r in range(10, 22) if isinstance(sk.cell(r, 2).value, (int, float))]
    sp = wv["SENSOR PT100"]
    sensor_tab = [{"indeks": float(sp.cell(r, 4).value), "koreksi": float(sp.cell(r, 5).value)}
                  for r in range(9, 16) if isinstance(sp.cell(r, 4).value, (int, float))]
    cmc = {
        "buret_digital": [{"maks_ml": 10, "nilai_ml": float(db["W6"].value)}, {"maks_ml": 50, "nilai_ml": float(db["W7"].value)}],
        "piston_pipette": [{"maks_ml": 1, "nilai_ml": float(db["Y6"].value)}, {"maks_ml": 5, "nilai_ml": float(db["Y7"].value)},
                           {"maks_ml": 10, "nilai_ml": float(db["Y8"].value)}],
        "dispensett": [{"maks_ml": 10, "nilai_ml": float(db["AA6"].value)}, {"maks_ml": 50, "nilai_ml": float(db["AA7"].value)},
                       {"maks_ml": 100, "nilai_ml": float(db["AA8"].value)}],
    }
    tm = wv["Tabel MPE"]
    mpe_pp = [{"nominal_ml": float(tm.cell(r, 3).value), "mpe_ul": float(tm.cell(r, 4).value)} for r in range(4, 17)]
    mpe_bd = [{"nominal_ml": float(tm.cell(r, 2).value), "hand_driven_ul": float(tm.cell(r, 3).value),
               "motor_driven_ul": float(tm.cell(r, 4).value)} for r in range(20, 28)]
    mpe_ds = []
    for r in range(31, 48):
        s = tm.cell(r, 3).value
        mpe_ds.append({"nominal_ml": float(tm.cell(r, 2).value),
                       "single_stroke_ul": float(s) if isinstance(s, (int, float)) else None,
                       "multi_stroke_ul": float(tm.cell(r, 4).value)})
    wv.close()
    return {
        "timbangan": timbangan, "termometer": termometer, "sensor": sensor,
        "koreksi_meter_suhu": meter, "koreksi_sensor_suhu": sensor_tab, "cmc": cmc,
        "mpe": {"piston_pipette": mpe_pp, "buret_digital": mpe_bd, "dispensett": mpe_ds},
    }


def timbangan(tabel, nama):
    for t in tabel["timbangan"]:
        if t["nama"] == nama:
            return t
    raise KeyError(nama)


def koreksi(tab, idx):
    for b in tab:
        if b["indeks"] == idx:
            return b["koreksi"]
    raise KeyError(idx)


def cmc_ml(tabel, jenis, nominal_ml, mode):
    """J48: pita master memakai batas BULAT (C33<=1, C33>=2..5) — celahnya 'cek range'."""
    pita = tabel["cmc"][jenis]
    if mode == "master":
        # Tiru persis: ml, batas bilangan bulat, celah -> None ("cek range")
        n = nominal_ml
        if jenis == "buret_digital":
            if n <= 10: return pita[0]["nilai_ml"]
            if 11 <= n <= 50: return pita[1]["nilai_ml"]
        elif jenis == "piston_pipette":
            if n <= 1: return pita[0]["nilai_ml"]
            if 2 <= n <= 5: return pita[1]["nilai_ml"]
            if 6 <= n <= 10: return pita[2]["nilai_ml"]
        else:
            if n <= 10: return pita[0]["nilai_ml"]
            if 11 <= n <= 50: return pita[1]["nilai_ml"]
            if 51 <= n <= 100: return pita[2]["nilai_ml"]
        return None
    for p in pita:                       # pita kontinu (G-7)
        if nominal_ml <= p["maks_ml"]:
            return p["nilai_ml"]
    return None


# ---------------------------------------------------------------------------
# Masukan
# ---------------------------------------------------------------------------
JENIS = {1: "buret_digital", 2: "dispensett", 3: "piston_pipette"}
SUB = {1: "single_stroke", 2: "multi_stroke", 3: "motor_driven", 4: "hand_driven"}


def masukan_fixed(path):
    _, wv = buka(path)
    i = wv["INPUT DATA"]
    m = {
        "keluarga": "fixed",
        "jenis": JENIS[i["E5"].value], "sub_jenis": SUB.get(i["E6"].value),
        "satuan": i["Y11"].value, "kapasitas": float(i["E15"].value),
        "timbangan": i["Y19"].value,
        "suhu_ruang": [float(i["E21"].value), float(i["F21"].value)],
        "kelembaban": [float(i["E22"].value), float(i["F22"].value)],
        "tekanan_udara": [float(i["E24"].value), float(i["F24"].value)],
        "titik": [{
            "label": "TUNGGAL", "nominal": float(i["E14"].value),
            "kumulatif": [float(i.cell(r, 8).value) for r in range(33, 44)],
            "suhu_air": [float(i["H49"].value), float(i["H50"].value)],
            "penguapan": 0.0,
        }],
    }
    wv.close()
    return m


def masukan_grad(path):
    _, wv = buka(path)
    i = wv["INPUT DATA"]
    pc = wv["PERHITUNGAN"]
    titik = []
    for lab, col in (("MIN", 8), ("MID", 10), ("MAX", 12)):
        titik.append({
            "label": lab, "nominal": float(i.cell(32, col).value),
            "kumulatif": [float(i.cell(r, col).value) for r in range(34, 45)],
            "suhu_air": [float(i.cell(50, col).value), float(i.cell(51, col).value)],
            "penguapan": float(pc.cell(52, col).value or 0.0),
        })
    m = {
        "keluarga": "graduated",
        "jenis": JENIS[i["E5"].value], "sub_jenis": SUB.get(i["E6"].value),
        "satuan": i["G15"].value, "kapasitas": float(i["E14"].value),
        "timbangan": i["Y19"].value,
        "suhu_ruang": [float(i["E22"].value), float(i["F22"].value)],
        "kelembaban": [float(i["E23"].value), float(i["F23"].value)],
        "tekanan_udara": [float(i["E25"].value), float(i["F25"].value)],
        "titik": titik,
    }
    wv.close()
    return m


# ---------------------------------------------------------------------------
# Rantai
# ---------------------------------------------------------------------------
def selisih(kum, mode, keluarga, label, satuan):
    """m_i = M_i - M_{i-1}; Graduated punya cabang tara-ulang di m4 & m7."""
    M = kum
    m = [M[k] - M[k - 1] for k in range(1, 11)]
    if keluarga == "graduated":
        tara = satuan == "ml"
        # m4 (indeks 3): IF(AND(G15="ml", M10>200), M4, M4-M3)
        if tara and M[10] > 200:
            m[3] = M[4]
        # m7 (indeks 6)
        if mode == "master":
            ambang = 20 if label == "MIN" else 200
            ambil = M[7] if label == "MIN" else M[4]        # G-8: J48/L48 menunjuk baris M4
        else:
            ambang, ambil = 200, M[7]
        if tara and M[10] > ambang:
            m[6] = ambil
    return m


def hitung(m, tabel, mode):
    kel = m["keluarga"]
    faktor_nominal = 0.001 if m["satuan"] in ("µl", "µL") else 1.0
    T = sum(m["suhu_ruang"]) / 2
    RH = sum(m["kelembaban"]) / 2
    P = sum(m["tekanan_udara"]) / 2
    ra = rho_udara(P, RH, T)
    idx_list = [b["indeks"] for b in tabel["koreksi_meter_suhu"]]

    per = []
    for t in m["titik"]:
        mi = selisih(t["kumulatif"], mode, kel, t["label"], m["satuan"])
        mi_k = [x + t["penguapan"] for x in mi]
        mbar = sum(mi_k) / len(mi_k)
        s = st.stdev(mi_k)
        ta, tb = t["suhu_air"]
        rata_baca = (ta + tb) / 2
        idx = terdekat(idx_list, rata_baca)
        km = koreksi(tabel["koreksi_meter_suhu"], idx)
        ks = koreksi(tabel["koreksi_sensor_suhu"], idx)
        t1 = ta + km + ks
        t2 = tb + km + ks
        tt = (t1 + t2) / 2
        r1, r2 = tanaka(t1), tanaka(t2)
        rw = (r1 + r2) / 2
        nominal_ml = t["nominal"] * faktor_nominal
        if kel == "fixed":
            V20 = mbar * (1 / (rw - ra)) * (1 - (ra / RHO_B)) * (1 - (ALFA * (tt - 20)))
        else:
            V20 = (mbar / RHO_B) * ((RHO_B - ra) / (rw - ra)) * (1 - (ALFA * (tt - 20)))
        per.append(dict(label=t["label"], nominal=t["nominal"], nominal_ml=nominal_ml, m=mi, m_terkoreksi=mi_k,
                        m_rata=mbar, stdev=s, suhu_rata_baca=rata_baca, indeks_suhu=idx, koreksi_meter=km,
                        koreksi_sensor=ks, t_terkoreksi=[t1, t2], t_rata=tt, rho_air_titik=[r1, r2], rho_air=rw,
                        V20=V20, deviasi=V20 - nominal_ml, deviasi_ul=(V20 - nominal_ml) * 1000))

    semua_t = [x for p in per for x in p["t_terkoreksi"]]
    semua_rho = [x for p in per for x in p["rho_air_titik"]]
    bal = timbangan(tabel, m["timbangan"])
    maks_s = max(p["stdev"] for p in per)

    if kel == "fixed":
        acuan = per[0]
        dT = max(semua_t) - min(semua_t)                  # J41-J42
        d_rho = max(semua_rho) - min(semua_rho)           # L40
        U_rho = math.sqrt(((1 / 1000000) / 2) ** 2 + d_rho ** 2)   # Q59
        div_rho, div_T = 2.0, math.sqrt(3)
        J3, J5, J7 = acuan["m_rata"], acuan["rho_air"], acuan["t_rata"]
    else:
        acuan_m = per[-1]                                  # J3 = L96 (titik MAX)
        acuan_r = per[0]                                   # J5 = H85, J7 = H67 (titik MIN)
        dT = max(semua_t) - min(semua_t)                  # N67-O67
        d_rho = max(semua_rho) - min(semua_rho)           # H86
        H87 = d_rho / 1.73
        P86 = 1 / 1000000
        U_rho = (math.sqrt(P86 ** 2) + (H87 ** 2)) if mode == "master" else math.sqrt(P86 ** 2 + H87 ** 2)   # G-9
        div_rho, div_T = math.sqrt(3), 2.0
        J3, J5, J7 = acuan_m["m_rata"], acuan_r["rho_air"], acuan_r["t_rata"]
    J4, J6, J8 = ra, RHO_B, ALFA

    res_bal = bal["resolusi"]
    F27 = res_bal / math.sqrt(3)
    F28 = maks_s / 2
    F29 = bal["u95"] / 2
    U_massa = math.sqrt(F27 ** 2 + F28 ** 2 + F29 ** 2)
    H15 = tabel["termometer"]["u95"] / 2
    H19 = tabel["sensor"]["u95"] / 2
    H23 = dT / (2 * math.sqrt(3))
    U_suhu = math.sqrt(H15 ** 2 + H19 ** 2 + H23 ** 2)

    muai = (1 - (J8 * (J7 - 20)))
    ci = [
        ((J6 - J4) / (J6 * (J5 - J4))) * muai,
        J3 * ((J6 - J5) / (J6 * ((J5 - J4) ** 2))) * muai,
        ((J5 - J4) - J3 * ((J6 - J4)) / (J6 * ((J5 - J4))) * 1) * muai if False else ((J5 - J4) - J3 * ((J6 - J4)) / (J6 * ((J5 - J4))) ) * muai,
        J3 * J4 / (J6 ** 2 * (J5 - J4)) * muai,
        (-J3 * J8 * (J6 - J4) * (J7 - 20) / (J6 * (J5 - J4))),
        -(J3 * (J6 - J4)) / (J6 * (J5 - J4)),
        1.0,
    ]
    komp_def = [
        ("massa_air", U_massa, 2.0, 18.0),
        ("densitas_udara", 0.1 * J4, math.sqrt(3), 50.0),
        ("densitas_air", U_rho, div_rho, 60.0),
        ("densitas_anak_timbangan", 0.1 * J6, math.sqrt(3), 50.0),
        ("suhu_air", U_suhu, div_T, 60.0),
        ("koefisien_muai", 0.1 * J8, math.sqrt(3), 40.0),
        ("operator", U_OPERATOR, 1.7320508075688772, 50.0),
    ]
    komp = []
    sJ = sK = 0.0
    for (kode, U, div, v), c in zip(komp_def, ci):
        G = U / div
        I = G * c
        J = I ** 2
        K = (J ** 2) / v
        sJ += J
        sK += K
        komp.append(dict(kode=kode, U=U, pembagi=div, vi=v, ci=c, u=G, uici=I, kuadrat=J, pangkat4_per_v=K))
    uc = math.sqrt(sJ)
    veff = (uc ** 4) / sK
    k = tinv(veff)
    Uexp = uc * k
    cmc = cmc_ml(tabel, m["jenis"], m["kapasitas"] * faktor_nominal, mode)
    if m["satuan"] in ("µl", "µL"):
        u95 = max(Uexp * 1000, (cmc or 0) * 1000) if cmc is not None else Uexp * 1000
    else:
        u95 = max(Uexp, cmc) if cmc is not None else Uexp
    return dict(per_titik=per, rho_udara=ra, suhu_ruang_rata=T, kelembaban_rata=RH, tekanan_rata=P,
                maks_stdev=maks_s, U_massa=U_massa, U_suhu=U_suhu, U_rho=U_rho, delta_suhu=dT, delta_rho=d_rho,
                acuan=dict(J3=J3, J4=J4, J5=J5, J6=J6, J7=J7, J8=J8),
                komponen=komp, jumlah_kuadrat=sJ, jumlah_pangkat4=sK, uc=uc, veff=veff, df=int(veff), k=k, U=Uexp,
                cmc=cmc, u95=u95)


# ---------------------------------------------------------------------------
hasil = []


def cek(label, harap, dapat, tol=TOL):
    if harap is None or dapat is None:
        ok = harap == dapat
        d = 0.0 if ok else float("nan")
    else:
        d = abs(float(harap) - float(dapat))
        ok = d <= tol * max(1.0, abs(float(harap)))
    hasil.append((label, harap, dapat, d, ok))


def adu_fixed(path, r):
    _, wv = buka(path)
    pc = wv["PERHITUNGAN"]; u = wv["PERHITUNGAN U95%"]; se = wv["SERTIFIKAT"]
    p = r["per_titik"][0]
    for k in range(10):
        cek(f"F m{k+1}", pc.cell(28 + k, 8).value, p["m"][k])
    cek("F H39 mbar", pc["H39"].value, p["m_rata"]); cek("F H40 s", pc["H40"].value, p["stdev"])
    cek("F H48 rata baca", pc["H48"].value, p["suhu_rata_baca"]); cek("F H49 idx", pc["H49"].value, p["indeks_suhu"])
    cek("F H50", pc["H50"].value, p["koreksi_meter"]); cek("F H51", pc["H51"].value, p["koreksi_sensor"])
    cek("F H52", pc["H52"].value, p["t_terkoreksi"][0]); cek("F H53", pc["H53"].value, p["t_terkoreksi"][1])
    cek("F H54", pc["H54"].value, p["t_rata"])
    cek("F L46", pc["L46"].value, p["rho_air_titik"][0]); cek("F L47", pc["L47"].value, p["rho_air_titik"][1])
    cek("F L48 rho_w", pc["L48"].value, p["rho_air"]); cek("F L40", pc["L40"].value, r["delta_rho"])
    cek("F H60 rho_a", pc["H60"].value, r["rho_udara"]); cek("F Q59", pc["Q59"].value, r["U_rho"])
    cek("F H63 V20", pc["H63"].value, p["V20"])
    cek("F U95 F30", u["F30"].value, r["U_massa"]); cek("F U95 H25", u["H25"].value, r["U_suhu"])
    for idx, row in enumerate(range(36, 43)):
        kk = r["komponen"][idx]
        cek(f"F U95 r{row} {kk['kode']} U", u.cell(row, 4).value, kk["U"])
        cek(f"F U95 r{row} div", u.cell(row, 5).value, kk["pembagi"])
        cek(f"F U95 r{row} vi", u.cell(row, 6).value, kk["vi"])
        cek(f"F U95 r{row} u", u.cell(row, 7).value, kk["u"])
        cek(f"F U95 r{row} ci", u.cell(row, 8).value, kk["ci"])
        cek(f"F U95 r{row} uici", u.cell(row, 9).value, kk["uici"])
        cek(f"F U95 r{row} ^2", u.cell(row, 10).value, kk["kuadrat"])
        cek(f"F U95 r{row} ^4/v", u.cell(row, 11).value, kk["pangkat4_per_v"])
    cek("F uc", u["I44"].value, r["uc"]); cek("F veff", u["I45"].value, r["veff"])
    cek("F k", u["I46"].value, r["k"]); cek("F U", u["I47"].value, r["U"])
    cek("F CMC", u["J48"].value, r["cmc"]); cek("F U95", u["J49"].value, r["u95"])
    cek("F sert N21 V20", se["N21"].value, p["V20"]); cek("F sert Q22", se["Q22"].value, r["u95"])
    cek("F sert V23 k", se["V23"].value, r["k"])
    wv.close()


def adu_grad(path, r):
    _, wv = buka(path)
    pc = wv["PERHITUNGAN"]; u = wv["PERHITUNGAN U95%"]; se = wv["SERTIFIKAT"]
    for ti, col in enumerate((8, 10, 12)):
        p = r["per_titik"][ti]
        L = p["label"]
        for k in range(10):
            cek(f"G {L} m{k+1}", pc.cell(42 + k, col).value, p["m"][k])
            cek(f"G {L} m'{k+1}", pc.cell(42 + k, col + 1).value, p["m_terkoreksi"][k])
        cek(f"G {L} mbar", pc.cell(53, col).value, p["m_rata"]); cek(f"G {L} s", pc.cell(54, col).value, p["stdev"])
        cek(f"G {L} rata baca", pc.cell(61, col).value, p["suhu_rata_baca"])
        cek(f"G {L} idx", pc.cell(62, col).value, p["indeks_suhu"])
        cek(f"G {L} km", pc.cell(63, col).value, p["koreksi_meter"]); cek(f"G {L} ks", pc.cell(64, col).value, p["koreksi_sensor"])
        cek(f"G {L} t1", pc.cell(65, col).value, p["t_terkoreksi"][0]); cek(f"G {L} t2", pc.cell(66, col).value, p["t_terkoreksi"][1])
        cek(f"G {L} trata", pc.cell(67, col).value, p["t_rata"])
        cek(f"G {L} r1", pc.cell(83, col).value, p["rho_air_titik"][0]); cek(f"G {L} r2", pc.cell(84, col).value, p["rho_air_titik"][1])
        cek(f"G {L} rw", pc.cell(85, col).value, p["rho_air"]); cek(f"G {L} ra", pc.cell(99, col).value, r["rho_udara"])
        cek(f"G {L} V20", pc.cell(102, col).value, p["V20"]); cek(f"G {L} dev", pc.cell(103, col).value, p["deviasi"])
        cek(f"G {L} dev ul", pc.cell(104, col).value, p["deviasi_ul"])
        cek(f"G sert N{21+ti}", se.cell(21 + ti, 14).value, p["V20"])
    cek("G O54 maks s", pc["O54"].value, r["maks_stdev"]); cek("G N67-O67", pc["N67"].value - pc["O67"].value, r["delta_suhu"])
    cek("G H86", pc["H86"].value, r["delta_rho"]); cek("G P87", pc["P87"].value, r["U_rho"])
    cek("G U95 F30", u["F30"].value, r["U_massa"]); cek("G U95 H25", u["H25"].value, r["U_suhu"])
    for idx, row in enumerate(range(36, 43)):
        kk = r["komponen"][idx]
        cek(f"G U95 r{row} {kk['kode']} U", u.cell(row, 4).value, kk["U"])
        cek(f"G U95 r{row} div", u.cell(row, 5).value, kk["pembagi"])
        cek(f"G U95 r{row} vi", u.cell(row, 6).value, kk["vi"])
        cek(f"G U95 r{row} u", u.cell(row, 7).value, kk["u"])
        cek(f"G U95 r{row} ci", u.cell(row, 8).value, kk["ci"])
        cek(f"G U95 r{row} uici", u.cell(row, 9).value, kk["uici"])
        cek(f"G U95 r{row} ^2", u.cell(row, 10).value, kk["kuadrat"])
        cek(f"G U95 r{row} ^4/v", u.cell(row, 11).value, kk["pangkat4_per_v"])
    cek("G uc", u["I44"].value, r["uc"]); cek("G veff", u["I45"].value, r["veff"])
    cek("G k", u["I46"].value, r["k"]); cek("G U", u["I47"].value, r["U"])
    cek("G CMC", u["J48"].value, r["cmc"]); cek("G U95", u["J49"].value, r["u95"])
    cek("G sert Q24", se["Q24"].value, r["u95"]); cek("G sert V25 k", se["V25"].value, r["k"])
    wv.close()


def sha256_berkas(path):
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for blok in iter(lambda: fh.read(1 << 20), b""):
            h.update(blok)
    return h.hexdigest()


def cek_manifest(akar, sha):
    """Workbook yang dibaca WAJIB versi yang dibekukan di manifest.

    Kalau tidak, JSON tidak ditulis: workbook baru = acuan baru, dan itu harus
    dicatat dulu (manifest + entri versi di log-metode-tekanan-piston.json)
    sebelum angka aplikasi boleh ikut bergeser.
    """
    with open(os.path.join(akar, "database/data/manifest-workbook-tekanan-piston.json"), encoding="utf-8") as fh:
        beku = {b["sha256"] for b in json.load(fh)["berkas"]}
    asing = {k: v for k, v in sha.items() if v not in beku}
    for k, v in asing.items():
        print(f"  TOLAK: {k} sha256 {v} tidak ada di manifest yang dibekukan")
    return not asing


def main():
    folder = sys.argv[1]
    tulis = sys.argv[3] if len(sys.argv) > 3 and sys.argv[2] == "--tulis" else None
    tf = baca_tabel(os.path.join(folder, FIXED))
    tg = baca_tabel(os.path.join(folder, GRAD))
    # Tabel referensi kedua master diadu: kalau beda, itu temuan, bukan gangguan.
    for kunci in ("timbangan", "koreksi_meter_suhu", "koreksi_sensor_suhu", "cmc", "mpe"):
        if tf[kunci] != tg[kunci]:
            print("  BEDA antar-master:", kunci)
    mf = masukan_fixed(os.path.join(folder, FIXED))
    mg = masukan_grad(os.path.join(folder, GRAD))
    rf_m = hitung(mf, tf, "master"); adu_fixed(os.path.join(folder, FIXED), rf_m)
    rg_m = hitung(mg, tg, "master"); adu_grad(os.path.join(folder, GRAD), rg_m)
    rf_b = hitung(mf, tf, "benar")
    rg_b = hitung(mg, tg, "benar")
    for nama, a, b in (("fixed", rf_m, rf_b), ("graduated", rg_m, rg_b)):
        print(f"{nama:10s} master uc={a['uc']!r} veff={a['veff']!r} U={a['U']!r} U95={a['u95']!r} | "
              f"benar U={b['U']!r} U95={b['u95']!r}")
    merah = [h for h in hasil if not h[4]]
    for h in merah:
        print("  MERAH", h[0], "excel=", repr(h[1]), "hitung=", repr(h[2]), "selisih=", h[3])
    terburuk = max((h[3] for h in hasil if h[3] == h[3]), default=0.0)
    print(f"\n{len(hasil) - len(merah)} hijau, {len(merah)} merah dari {len(hasil)} sel; selisih terbesar {terburuk:.3e}")

    sha = {"fixed": sha256_berkas(os.path.join(folder, FIXED)), "graduated": sha256_berkas(os.path.join(folder, GRAD))}
    if tulis and not merah and not cek_manifest(tulis, sha):
        sys.exit(2)
    if tulis and not merah:
        sumber = {
            "dibuat_oleh": "docs/skrip/gen-tabel-standar-piston-volume.py",
            "tanggal": datetime.date.today().isoformat(),
            "workbook": NAMA_ASLI,
            "workbook_sha256": sha,
            "catatan": "Digenerate dari dua Master Olah Data piston volume. JANGAN disunting tangan. "
                       "Identitas pelanggan tidak ikut diekspor.",
        }
        tabel = {
            "_sumber": sumber,
            "konstanta": {"tanaka": {"a1": A1, "a2": A2, "a3": A3, "a4": A4, "a5": A5},
                          "densitas_anak_timbangan": RHO_B, "koefisien_muai": ALFA, "u_operator": U_OPERATOR},
            **tf,
        }
        with open(os.path.join(tulis, "database/data/tabel-standar-piston-volume.json"), "w", encoding="utf-8", newline="\n") as fh:
            json.dump(tabel, fh, ensure_ascii=False, indent=2); fh.write("\n")
        sesi = {"_sumber": sumber, "sesi": {
            "fixed": {"masukan": mf, "harapan_master_excel": rf_m, "harapan_benar": rf_b},
            "graduated": {"masukan": mg, "harapan_master_excel": rg_m, "harapan_benar": rg_b},
        }}
        with open(os.path.join(tulis, "database/data/sesi-master-piston-volume.json"), "w", encoding="utf-8", newline="\n") as fh:
            json.dump(sesi, fh, ensure_ascii=False, indent=2); fh.write("\n")
        print("JSON ditulis.")
    sys.exit(1 if merah else 0)


if __name__ == "__main__":
    main()
