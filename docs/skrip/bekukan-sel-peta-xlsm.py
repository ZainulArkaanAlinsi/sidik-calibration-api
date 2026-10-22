#!/usr/bin/env python3
"""Bekukan SEL YANG DIRUJUK peta sel dari master .xlsm asli jadi fixture test.

    python -I docs/skrip/bekukan-sel-peta-xlsm.py <akar-repo> <paket> <berkas.xlsm>...

Untuk paket yang TIDAK punya ekspor CSV di repo (Tekanan, Piston Volume, ...),
paritas PHP (`PemeriksaMaster`) lawan JSON repo butuh sel master. Master tidak
boleh masuk repo (data pelanggan), jadi yang dibekukan HANYA sel yang dirujuk
peta: jangkar, semua rentang `ekstrak` (termasuk `ulang`), sel `cek`, dan
rentang gabungan (merge) yang menyentuhnya beserta sel kiri-atasnya. Isinya
data acuan yang memang sudah ada di JSON repo + label jangkar.

Keluaran: `tests/Fixtures/peta-sel-xlsm/<paket>/<varian|nama>.json` berbentuk
keluaran `PembacaXlsmWorkbook` (`lembar[].nama`, `sel_kanonik`, `tata_letak.merge`)
supaya `PemeriksaMaster::periksa(.., FORMAT_XLSM)` membacanya apa adanya.
Bentuk sel meniru pembaca PHP: angka `t=n`, teks `t=s`, galat `t=e`, rumus
`t=f` (+`r`), tanggal = nomor seri Excel di `n`. `s` angka = `repr` Python,
BUKAN teks berformat Excel (pembaca PHP memformatnya) — peta tidak memakai
teks tampil sel angka kecuali bilangan bulat, yang dibaca dari `n`.

Sandi dari `SIDIK_MASTER_SANDI` (HKCU Environment) seperti verifikasi-peta-xlsm.py;
dekripsi ke memori saja.
"""
import datetime
import hashlib
import io
import json
import os
import re
import sys
import winreg

import msoffcrypto
import openpyxl
from openpyxl.utils import get_column_letter

REPO, PAKET = sys.argv[1], sys.argv[2]
BERKAS = sys.argv[3:]


def sandi():
    with winreg.OpenKey(winreg.HKEY_CURRENT_USER, 'Environment') as k:
        return winreg.QueryValueEx(k, 'SIDIK_MASTER_SANDI')[0]


def buka(path):
    data = open(path, 'rb').read()
    buf = io.BytesIO(data)
    if data[:4] == b'\xd0\xcf\x11\xe0':
        out = io.BytesIO()
        f = msoffcrypto.OfficeFile(buf)
        f.load_key(password=sandi())
        f.decrypt(out)
        buf = out
    buf.seek(0)
    nilai = openpyxl.load_workbook(io.BytesIO(buf.getvalue()), data_only=True)
    rumus = openpyxl.load_workbook(io.BytesIO(buf.getvalue()), data_only=False)
    return nilai, rumus, hashlib.sha256(data).hexdigest()


def kol(h):
    n = 0
    for c in h:
        n = n * 26 + ord(c) - 64
    return n


def urai(a):
    i = a.rindex('!')
    m = re.match(r"^([A-Z]+)(\d+)(?::([A-Z]+)(\d+))?$", a[i + 1:])
    k1, b1 = kol(m[1]), int(m[2])
    k2, b2 = (kol(m[3]), int(m[4])) if m[3] else (k1, b1)
    return a[:i], k1, b1, k2, b2


def normal_nama(n):
    n = re.sub(r"\.(xlsm|xlsx)$", "", n, flags=re.I)
    return re.sub(r"[^a-z0-9.]+", " ", n.lower()).strip()


def berlaku(b, var):
    return 'varian' not in b or (var is not None and var in b['varian'])


peta = json.load(open(f"{REPO}/database/data/peta-sel/{PAKET}.json", encoding='utf-8'))
x = peta.get('xlsm') or {}
NAMA_ASLI = dict(x.get('lembar', {}))

for path in BERKAS:
    nama = os.path.basename(path)
    var = None
    if 'varian' in peta:
        cocok = [k for k, rx in peta['varian']['pola'].items() if re.search(rx, normal_nama(nama))]
        if len(cocok) != 1:
            sys.exit(f"{nama}: varian tidak dikenali")
        var = cocok[0]
    geser = {}
    for g in x.get('geser', []):
        if berlaku(g, var):
            geser.setdefault(g['lembar'], []).append((g['dari'], g['geser']))

    def baris_asli(lb, b):
        gs = 0
        for dari, n in sorted(geser.get(lb, [])):
            if b >= dari:
                gs = n
        return b + gs

    rujuk = {}  # lembar peta -> set (kolom, baris asli)

    def tambah(alamat, langkah=0):
        lb, k1, b1, k2, b2 = urai(alamat)
        for b in range(b1, b2 + 1):
            for k in range(k1, k2 + 1):
                rujuk.setdefault(lb, set()).add((k, baris_asli(lb, b + langkah)))

    for j in peta['jangkar']:
        if berlaku(j, var):
            tambah(j['sel'])
    for daftar in (peta['ekstrak'], peta.get('cek', [])):
        for e in daftar:
            if not berlaku(e, var):
                continue
            for i in range(e.get('ulang') or 1):
                langkah = i * e.get('langkah_baris', 0)
                for kunci in ('sel', 'rentang', 'kunci', 'nilai'):
                    for alamat in (e[kunci] if isinstance(e.get(kunci), list) else [e.get(kunci)]):
                        if isinstance(alamat, str) and '!' in alamat:
                            tambah(alamat, langkah)
                for sp in (e.get('kolom') or {}).values():
                    r = sp.get('rentang') if isinstance(sp, dict) else sp
                    if isinstance(r, str):
                        tambah(r, langkah)

    wv, wf, sha = buka(path)
    lembar = []
    for lb, sel in sorted(rujuk.items()):
        asli = NAMA_ASLI.get(lb, lb)
        ws, wr = wv[asli], wf[asli]
        merge = []
        for rg in ws.merged_cells.ranges:
            if any(rg.min_col <= k <= rg.max_col and rg.min_row <= b <= rg.max_row for k, b in sel):
                merge.append(str(rg))
                sel.add((rg.min_col, rg.min_row))
        hasil = []
        for k, b in sorted(sel, key=lambda kb: (kb[1], kb[0])):
            v, f = ws.cell(row=b, column=k).value, wr.cell(row=b, column=k).value
            if v is None and not (isinstance(f, str) and f.startswith('=')):
                continue
            a = f"{get_column_letter(k)}{b}"
            c = {'a': a, 'b': b, 'k': k, 't': 's', 'n': None, 's': '', 'r': None, 'f': None, 'g': None}
            if isinstance(v, bool):
                c.update(t='b', s='TRUE' if v else 'FALSE')
            elif isinstance(v, (datetime.datetime, datetime.date)):
                d = v if isinstance(v, datetime.datetime) else datetime.datetime(v.year, v.month, v.day)
                serial = (d - datetime.datetime(1899, 12, 30)).total_seconds() / 86400
                c.update(t='n', n=serial, s=repr(serial))
            elif isinstance(v, (int, float)):
                c.update(t='n', n=float(v), s=repr(v))
            elif isinstance(v, str) and v.startswith('#'):
                c.update(t='e', s=v)
            elif v is not None:
                c.update(s=str(v))
            if isinstance(f, str) and f.startswith('='):
                c.update(t='f', r=f)
            hasil.append(c)
        lembar.append({'nama': asli, 'sel_kanonik': json.dumps(hasil, ensure_ascii=False), 'tata_letak': {'merge': sorted(merge)}})
    keluar = f"{REPO}/tests/Fixtures/peta-sel-xlsm/{PAKET}"
    os.makedirs(keluar, exist_ok=True)
    berkas = f"{keluar}/{var or normal_nama(nama).replace(' ', '_')}.json"
    with open(berkas, 'w', encoding='utf-8', newline='\n') as fh:
        json.dump({'_sumber': 'docs/skrip/bekukan-sel-peta-xlsm.py — HANYA sel yang dirujuk peta', 'berkas': nama, 'sha256': sha,
                   'varian': var, 'lembar': lembar}, fh, ensure_ascii=False, indent=1)
        fh.write('\n')
    print(berkas, sum(len(json.loads(l['sel_kanonik'])) for l in lembar), 'sel')
